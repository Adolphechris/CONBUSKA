from django.core.management.base import BaseCommand

from fournisseurs.models import Fournisseur
from produits.models import Article, Categorie, Unite
from openpyxl import load_workbook
from django.db.utils import IntegrityError


class Command(BaseCommand):
    help = "Create new articles through excel file"

    def handle(self, *args, **options):
        filename = 'produits/management/commands/articles.xlsx'

        wb = load_workbook(filename=filename)
        print(wb.sheetnames)
        sheet = wb['articles']
        header = ['id', 'nom', 'est', 'embalage', 'price', 'stock', 'stockDepot', 'priceGros', 'seuilRouge', 'QteGros',
                  'idprice', 'creatAt', 'deleteAt', 'updateAt']

        rows = []
        for row in sheet.iter_rows():
            r = dict()
            for h, cell in zip(header, row):
                r[h] = cell.value

            if r['nom'] is not None:
                rows.append(r)

        for i in rows[1::]:
            if i.get('nom') is not None:
                print(i.get('nom'))
                if i.get('deleteAt') is None:
                    categorie = Categorie.objects.get(nom='All')
                    unite = Unite.objects.get(nom='Pce')
                    fournisseur = Fournisseur.objects.get(nom='Hyper Psaro')

                    try:
                        Article.objects.create(
                            designation=i.get('nom'),
                            description="",
                            code_barre="",
                            categorie=categorie,
                            unite=unite,
                            fournisseur=fournisseur,
                            prix_achat=i.get('price'),
                            prix_vente=i.get('price'),
                            prix_vente_gros=i.get('priceGros'),
                            devise='FC',
                            seuil=i.get('seuilRouge'),
                            seuil_gros=i.get('QteGros'),
                            emplacement='Alimentation'
                        )
                    except IntegrityError:
                        print('*Doublon: Ce produit est deja enregistre')
                else:
                    print('## Article supprimé')